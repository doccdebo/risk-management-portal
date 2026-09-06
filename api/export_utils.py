import csv
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from jinja2 import Environment, FileSystemLoader, select_autoescape
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, A4
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

COLUMNS = [
    ("uid", "ID"),
    ("title", "Title"),
    ("state", "State"),
    ("assigned_to", "Assigned To"),
    ("likelihood", "Likelihood"),
    ("impact", "Impact"),
    ("initial_risk", "Initial Risk"),
    ("threat_type", "Threat Type"),
    ("threat_score", "Threat Score"),
    ("threat_severity", "Threat Severity"),
    ("risk_sub_type", "Risk SubType"),
    ("service", "Service"),
    ("build_number", "Build Number"),
    ("mitigation_plan", "Mitigation Plan"),
    ("contingency_plan", "Contingency Plan"),
    ("created_by", "Created By"),
    ("created_date", "Created Date"),
    ("changed_date", "Changed Date"),
]

_env = Environment(
    loader=FileSystemLoader("templates"),
    autoescape=select_autoescape(["html"]),
)


def _row_value(risk: dict, key: str):
    value = risk.get(key, "")
    if value is None:
        return ""
    return value


def to_csv(risks: list[dict]) -> str:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([label for _, label in COLUMNS])
    for risk in risks:
        writer.writerow([_row_value(risk, key) for key, _ in COLUMNS])
    return buffer.getvalue()


def to_excel_bytes(risks: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Risks"

    header_fill = PatternFill(start_color="0078D4", end_color="0078D4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for col_idx, (_, label) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.fill = header_fill
        cell.font = header_font

    for row_idx, risk in enumerate(risks, start=2):
        for col_idx, (key, _) in enumerate(COLUMNS, start=1):
            ws.cell(row=row_idx, column=col_idx, value=_row_value(risk, key))

    for col_idx, (key, label) in enumerate(COLUMNS, start=1):
        max_len = max([len(label)] + [len(str(_row_value(r, key))) for r in risks]) if risks else len(label)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 40)

    ws.freeze_panes = "A2"

    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


def to_html(risks: list[dict], title: str = "Risks") -> str:
    template = _env.get_template("risks_export.html")
    return template.render(title=title, risks=risks, columns=COLUMNS)


def to_pdf_bytes(risks: list[dict], title: str = "Risks") -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        topMargin=14 * mm, bottomMargin=10 * mm, leftMargin=10 * mm, rightMargin=10 * mm,
    )
    styles = getSampleStyleSheet()
    title_style = styles["Heading1"]
    title_style.textColor = colors.HexColor("#0078D4")

    print_cols = ["uid", "title", "state", "assigned_to", "likelihood", "impact", "initial_risk", "threat_score", "threat_severity"]
    labels = [label for key, label in COLUMNS if key in print_cols]

    data = [labels]
    for risk in risks:
        data.append([str(_row_value(risk, key)) for key in print_cols])

    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0078D4")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D0D0D0")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F6F8FA")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))

    elements = [
        Paragraph(title, title_style),
        Paragraph(f"Exported on {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", styles["Normal"]),
        Spacer(1, 8),
        table,
    ]
    doc.build(elements)
    return buffer.getvalue()
